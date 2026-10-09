"""Tạo bộ ví dụ phiếu lương: danh sách Excel + file thành phẩm mẫu Word (+ Excel) có ký hiệu [Tên cột]."""
import datetime, openpyxl, sys
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = sys.argv[1] if len(sys.argv) > 1 else '.'
KY = '10/2026'
CTY = 'CÔNG TY TNHH ABC VIỆT NAM'

# ---------------- số thành chữ ----------------
DIG = ['không', 'một', 'hai', 'ba', 'bốn', 'năm', 'sáu', 'bảy', 'tám', 'chín']
def doc3(n, full):
    tr, ch, dv = n // 100, n // 10 % 10, n % 10
    s = []
    if full or tr: s += [DIG[tr], 'trăm']
    if ch == 0:
        if dv and (full or tr): s.append('linh')
    elif ch == 1: s.append('mười')
    else: s += [DIG[ch], 'mươi']
    if dv:
        if dv == 1 and ch > 1: s.append('mốt')
        elif dv == 5 and ch > 0: s.append('lăm')
        elif dv == 4 and ch > 1: s.append('tư')
        else: s.append(DIG[dv])
    return s
def bang_chu(n):
    if n == 0: return 'Không đồng'
    units = ['', 'nghìn', 'triệu', 'tỷ']
    parts, i = [], 0
    while n: parts.append(n % 1000); n //= 1000
    out = []
    for i in range(len(parts) - 1, -1, -1):
        g = parts[i]
        if g == 0: continue
        out += doc3(g, i < len(parts) - 1) + ([units[i]] if units[i] else [])
    s = ' '.join(out) + ' đồng'
    return s[0].upper() + s[1:]

# ---------------- dữ liệu minh hoạ ----------------
NV = [  # mã, họ tên, phòng ban, chức vụ, công, phép, lương CB, hiệu quả, OT, điện thoại, công tác, tạm ứng, người phụ thuộc, email
    ('NV001', 'Nguyễn Văn An', 'Kinh doanh', 'Trưởng phòng', 22, 0, 25000000, 6000000, 1500000, 500000, 1200000, 5000000, 1, 'an.nv@abc.vn'),
    ('NV002', 'Trần Thị Bích Ngọc', 'Hành chính – Nhân sự', 'Chuyên viên', 21, 1, 14500000, 2500000, 0, 300000, 0, 3000000, 0, 'ngoc.ttb@abc.vn'),
    ('NV003', 'Lê Hoàng Minh', 'Kế toán', 'Kế toán trưởng', 22, 0, 22000000, 4000000, 800000, 500000, 0, 0, 2, 'minh.lh@abc.vn'),
    ('NV004', 'Phạm Thu Hà', 'Kinh doanh', 'Nhân viên', 20, 2, 11000000, 3200000, 600000, 300000, 900000, 2000000, 0, 'ha.pt@abc.vn'),
    ('NV005', 'Võ Quốc Bảo', 'Kỹ thuật', 'Kỹ sư', 22, 0, 18750000, 3000000, 2400000, 300000, 1500000, 0, 1, 'bao.vq@abc.vn'),
    ('NV006', 'Đặng Ngọc Ánh', 'Chăm sóc khách hàng', 'Nhân viên', 21, 1, 9800000, 1500000, 450000, 200000, 0, 1500000, 0, 'anh.dn@abc.vn'),
    ('NV007', 'Huỳnh Đức Thịnh', 'Ban giám đốc', 'Giám đốc chi nhánh', 22, 0, 45000000, 10000000, 0, 1000000, 2500000, 10000000, 2, 'thinh.hd@abc.vn'),
    ('NV008', 'Bùi Thị Mai Phương', 'Marketing', 'Chuyên viên', 22, 0, 13200000, 2200000, 350000, 300000, 600000, 0, 0, 'phuong.btm@abc.vn'),
]
AN_CA = 730000
def thue(tn):
    """Biểu luỹ tiến minh hoạ 5 bậc (5/10/20/30/35%) – chỉ để làm ví dụ."""
    if tn <= 0: return 0
    b = [(10e6, .05), (30e6, .10), (60e6, .20), (100e6, .30), (float('inf'), .35)]
    t, lo = 0, 0
    for hi, r in b:
        if tn > lo: t += (min(tn, hi) - lo) * r
        lo = hi
    return int(round(t, -3))

HEAD = ['STT', 'Mã NV', 'Họ và tên', 'Phòng ban', 'Chức vụ', 'Kỳ lương', 'Ngày công', 'Nghỉ phép',
        'Lương cơ bản', 'Lương hiệu quả', 'Làm thêm giờ', 'Phụ cấp ăn ca', 'Phụ cấp điện thoại', 'Công tác phí', 'Tổng thu nhập',
        'BHXH', 'BHYT', 'BHTN', 'Thuế TNCN', 'Tạm ứng', 'Tổng khấu trừ', 'Thực lĩnh', 'Bằng chữ', 'Ghi chú', 'Email']
MONEY = set(HEAD[8:22])
NOTES = {'NV001': 'Thưởng doanh số quý III đã gộp vào lương hiệu quả', 'NV004': 'Nghỉ phép 2 ngày (14–15/10)', 'NV007': 'Tạm ứng công tác Đà Nẵng', 'NV005': 'Làm thêm giờ dự án lắp đặt kho Bắc Ninh'}
rows = []
for i, (ma, ten, pb, cv, cong, phep, cb, hq, ot, dt, ctp, tu, npt, mail) in enumerate(NV, 1):
    tong = cb + hq + ot + AN_CA + dt + ctp
    bhxh, bhyt, bhtn = round(cb * .08), round(cb * .015), round(cb * .01)
    tnct = tong - AN_CA - ctp - bhxh - bhyt - bhtn - 15500000 - npt * 6200000
    tax = thue(tnct)
    kt = bhxh + bhyt + bhtn + tax + tu
    tl = tong - kt
    rows.append([i, ma, ten, pb, cv, KY, cong, phep, cb, hq, ot, AN_CA, dt, ctp, tong, bhxh, bhyt, bhtn, tax, tu, kt, tl, bang_chu(tl), NOTES.get(ma, 'Không'), mail])

# ---------------- danh sách Excel ----------------
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Danh sách lương'
ws.append(HEAD)
for r in rows: ws.append(r)
hf = PatternFill('solid', fgColor='1F4E5F'); thin = Side(style='thin', color='C9D3D8')
for j, hd in enumerate(HEAD, 1):
    c = ws.cell(1, j); c.font = Font(bold=True, color='FFFFFF', name='Arial', size=10); c.fill = hf
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    w = 34 if hd in ('Bằng chữ', 'Ghi chú') else 22 if hd in ('Họ và tên', 'Phòng ban', 'Email') else 16 if hd in ('Chức vụ',) else 14 if hd in MONEY else 10
    ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = w
    for i in range(2, len(rows) + 2):
        cc = ws.cell(i, j); cc.font = Font(name='Arial', size=10); cc.border = Border(bottom=thin)
        if hd in MONEY: cc.number_format = '#,##0'
ws.row_dimensions[1].height = 32; ws.freeze_panes = 'D2'
wb.save(f'{OUT}/danh-sach-luong-mau.xlsx')

# ---------------- Word: phiếu lương ----------------
TEAL = RGBColor(0x1F, 0x4E, 0x5F); MUTE = RGBColor(0x6B, 0x77, 0x7D)
def shade(cell, hexc):
    tcPr = cell._tc.get_or_add_tcPr(); s = OxmlElement('w:shd'); s.set(qn('w:val'), 'clear'); s.set(qn('w:color'), 'auto'); s.set(qn('w:fill'), hexc); tcPr.append(s)
def borders(tbl, color='C9D3D8', inside=True, sz=4):
    tblPr = tbl._tbl.tblPr; b = OxmlElement('w:tblBorders')
    for e in ['top', 'left', 'bottom', 'right'] + (['insideH', 'insideV'] if inside else []):
        x = OxmlElement(f'w:{e}'); x.set(qn('w:val'), 'single'); x.set(qn('w:sz'), str(sz)); x.set(qn('w:color'), color); b.append(x)
    for e in ([] if inside else ['insideH', 'insideV']):
        x = OxmlElement(f'w:{e}'); x.set(qn('w:val'), 'nil'); b.append(x)
    tblPr.append(b)
def no_borders(tbl):
    tblPr = tbl._tbl.tblPr; b = OxmlElement('w:tblBorders')
    for e in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        x = OxmlElement(f'w:{e}'); x.set(qn('w:val'), 'nil'); b.append(x)
    tblPr.append(b)
def widths(tbl, cms):
    tbl.autofit = False
    for g, w in zip(tbl._tbl.tblGrid.findall(qn('w:gridCol')), cms): g.set(qn('w:w'), str(int(w * 567)))
    tblPr = tbl._tbl.tblPr; lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
    tw = OxmlElement('w:tblW'); tw.set(qn('w:w'), str(int(sum(cms) * 567))); tw.set(qn('w:type'), 'dxa')
    for x in tblPr.findall(qn('w:tblW')): tblPr.remove(x)
    tblPr.append(tw)
    for row in tbl.rows:
        cells = row._tr.findall(qn('w:tc')); i = 0
        for tc in cells:
            span = tc.tcPr.find(qn('w:gridSpan')) if tc.tcPr is not None else None
            n = int(span.get(qn('w:val'))) if span is not None else 1
            w = sum(cms[i:i + n]); i += n
            tcPr = tc.get_or_add_tcPr(); [tcPr.remove(x) for x in tcPr.findall(qn('w:tcW'))]
            x = OxmlElement('w:tcW'); x.set(qn('w:w'), str(int(w * 567))); x.set(qn('w:type'), 'dxa'); tcPr.insert(0, x)
def cell_pad(tbl, top=60, bottom=60):
    tblPr = tbl._tbl.tblPr; m = OxmlElement('w:tblCellMar')
    for e, v in [('top', top), ('bottom', bottom), ('left', 110), ('right', 110)]:
        x = OxmlElement(f'w:{e}'); x.set(qn('w:w'), str(v)); x.set(qn('w:type'), 'dxa'); m.append(x)
    tblPr.append(m)
def put(cell, text, b=False, i=False, sz=11, color=None, align=None, after=0):
    p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(after); p.paragraph_format.space_before = Pt(0)
    if align: p.alignment = align
    add(p, text, b, i, sz, color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    return p
def add(p, text, b=False, i=False, sz=11, color=None):
    r = p.add_run(text); r.bold = b; r.italic = i; r.font.size = Pt(sz)
    if color: r.font.color.rgb = color
    return r
def para(doc, text='', b=False, i=False, sz=11, color=None, align=None, before=0, after=4):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(after)
    if align is not None: p.alignment = align
    if text: add(p, text, b, i, sz, color)
    return p

doc = Document()
st = doc.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
sec = doc.sections[0]; sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2); sec.top_margin = Cm(1.6); sec.bottom_margin = Cm(1.5)
R, C, L = WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT

# đầu phiếu
t = doc.add_table(rows=1, cols=2); no_borders(t); widths(t, [11.5, 5.5])
c0, c1 = t.rows[0].cells
put(c0, CTY, b=True, sz=12, color=TEAL)
p = c0.add_paragraph(); add(p, 'Tầng 5, Toà nhà Mẫu, Q. Cầu Giấy, Hà Nội · MST 0100000000', sz=9, color=MUTE)
put(c1, 'Mã nhân viên', sz=9, color=MUTE, align=R)
p = c1.add_paragraph(); p.alignment = R; add(p, '[Mã NV]', b=True, sz=12, color=TEAL)
# vạch kẻ
p = para(doc, after=2); pPr = p._p.get_or_add_pPr(); bd = OxmlElement('w:pBdr'); bt = OxmlElement('w:bottom')
for k, v in [('val', 'single'), ('sz', '12'), ('space', '1'), ('color', '1F4E5F')]: bt.set(qn('w:' + k), v)
bd.append(bt); pPr.append(bd)
para(doc, 'PHIẾU LƯƠNG', b=True, sz=22, color=TEAL, align=C, before=6, after=0)
p = para(doc, align=C, after=12); add(p, 'Tháng ', i=True, sz=12, color=MUTE); add(p, '[Kỳ lương]', b=True, i=True, sz=12, color=MUTE)

# thông tin nhân viên
info = [('Họ và tên', '[Họ và tên]', 'Phòng ban', '[Phòng ban]'), ('Chức vụ', '[Chức vụ]', 'Ngày công thực tế', '[Ngày công] ngày'), ('Email', '[Email]', 'Nghỉ phép', '[Nghỉ phép] ngày')]
t = doc.add_table(rows=len(info), cols=4); borders(t, 'D5DEE2', inside=True); cell_pad(t, 45, 45); widths(t, [3.0, 5.6, 3.6, 4.8])
for r, (a, b, c, d) in zip(t.rows, info):
    for k, (txt, isv) in enumerate([(a, 0), (b, 1), (c, 0), (d, 1)]):
        cl = r.cells[k]
        if not isv: shade(cl, 'EEF3F5'); put(cl, txt, sz=10, color=MUTE)
        else: put(cl, txt, b=(txt == '[Họ và tên]'), sz=11.5 if txt == '[Họ và tên]' else 11)
para(doc, after=6)

# bảng thu nhập / khấu trừ
items = [('sec', 'I. THU NHẬP'), ('1', 'Lương cơ bản'), ('2', 'Lương hiệu quả'), ('3', 'Làm thêm giờ'), ('4', 'Phụ cấp ăn ca'), ('5', 'Phụ cấp điện thoại'), ('6', 'Công tác phí'),
         ('tot', 'Tổng thu nhập'),
         ('sec', 'II. CÁC KHOẢN KHẤU TRỪ'), ('1', 'BHXH', 'Bảo hiểm xã hội (8%)'), ('2', 'BHYT', 'Bảo hiểm y tế (1,5%)'), ('3', 'BHTN', 'Bảo hiểm thất nghiệp (1%)'), ('4', 'Thuế TNCN', 'Thuế thu nhập cá nhân'), ('5', 'Tạm ứng', 'Tạm ứng trong kỳ'),
         ('tot', 'Tổng khấu trừ')]
t = doc.add_table(rows=1, cols=3); borders(t, 'D5DEE2', inside=True); cell_pad(t, 30, 30)
hd = t.rows[0].cells
for cl, txt, al in zip(hd, ['STT', 'Khoản mục', 'Số tiền (đồng)'], [C, L, R]): shade(cl, '1F4E5F'); put(cl, txt, b=True, sz=10, color=RGBColor(255, 255, 255), align=al)
for it in items:
    r = t.add_row().cells
    if it[0] == 'sec':
        m = r[0].merge(r[2]); shade(m, 'E3ECEF'); put(m, it[1], b=True, sz=10.5, color=TEAL)
    elif it[0] == 'tot':
        m = r[0].merge(r[1]); shade(m, 'F6F8F9'); shade(r[2], 'F6F8F9')
        put(m, it[1], b=True, align=R); put(r[2], f'[{it[1]}]', b=True, align=R)
    else:
        col, lab = (it[1], it[2]) if len(it) == 3 else (it[1], it[1])
        put(r[0], it[0], sz=10, color=MUTE, align=C); put(r[1], lab); put(r[2], f'[{col}]', align=R)
widths(t, [1.3, 10.2, 5.5])

# thực lĩnh
para(doc, after=6)
t = doc.add_table(rows=1, cols=2); no_borders(t); cell_pad(t, 90, 90); widths(t, [8.5, 8.5])
a, b = t.rows[0].cells; shade(a, '1F4E5F'); shade(b, '1F4E5F')
put(a, 'THỰC LĨNH', b=True, sz=13, color=RGBColor(255, 255, 255))
p = put(b, '[Thực lĩnh]', b=True, sz=16, color=RGBColor(0xFF, 0xE2, 0xA8), align=R); add(p, ' đồng', sz=11, color=RGBColor(255, 255, 255))
p = para(doc, before=6, after=2); add(p, 'Bằng chữ: ', i=True, sz=10.5, color=MUTE); add(p, '[Bằng chữ]', i=True, sz=10.5)
p = para(doc, after=14); add(p, 'Ghi chú: ', i=True, sz=10.5, color=MUTE); add(p, '[Ghi chú]', i=True, sz=10.5)

# ký
t = doc.add_table(rows=1, cols=2); no_borders(t); widths(t, [8.5, 8.5])
for cl, ttl, nm in zip(t.rows[0].cells, ['NGƯỜI LẬP PHIẾU', 'NGƯỜI NHẬN LƯƠNG'], ['Nguyễn Thị Lan', '[Họ và tên]']):
    put(cl, ttl, b=True, sz=10.5, align=C)
    p = cl.add_paragraph(); p.alignment = C; add(p, '(Ký, ghi rõ họ tên)', i=True, sz=9.5, color=MUTE)
    for _ in range(2): cl.add_paragraph()
    p = cl.add_paragraph(); p.alignment = C; add(p, nm, b=True, sz=11)
p = para(doc, before=10, align=C, after=0)
add(p, 'Phiếu lương là thông tin cá nhân, vui lòng không chia sẻ. Mọi thắc mắc xin liên hệ phòng Hành chính – Nhân sự trong 3 ngày làm việc.', i=True, sz=8.5, color=MUTE)
doc.core_properties.author = 'Quản trị tử tế'; doc.core_properties.title = 'Phiếu lương mẫu'
doc.save(f'{OUT}/mau-phieu-luong.docx')

# ---------------- Excel: phiếu lương (mẫu Excel) ----------------
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Phiếu lương'
ws.sheet_view.showGridLines = False
for col, w in zip('ABC', [6, 34, 20]): ws.column_dimensions[col].width = w
tf = Font(name='Times New Roman', bold=True, size=16, color='1F4E5F'); nf = Font(name='Times New Roman', size=11); bf = Font(name='Times New Roman', size=11, bold=True)
mf = Font(name='Times New Roman', size=10, color='6B777D'); line = Side(style='thin', color='D5DEE2')
ws['A1'] = CTY; ws['A1'].font = Font(name='Times New Roman', bold=True, size=11, color='1F4E5F')
ws['A3'] = 'PHIẾU LƯƠNG THÁNG [Kỳ lương]'; ws['A3'].font = tf; ws.merge_cells('A3:C3'); ws['A3'].alignment = Alignment(horizontal='center')
r = 5
for lab, v in [('Họ và tên', '[Họ và tên]'), ('Mã NV', '[Mã NV]'), ('Phòng ban', '[Phòng ban]'), ('Chức vụ', '[Chức vụ]'), ('Ngày công', '[Ngày công]')]:
    ws.cell(r, 1, lab).font = mf; ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2); ws.cell(r, 3, v).font = bf if lab == 'Họ và tên' else nf; r += 1
r += 1
def money_row(lab, key, bold=False, fill=None):
    global r
    ws.cell(r, 2, lab).font = bf if bold else nf
    c = ws.cell(r, 3, key); c.font = bf if bold else nf; c.number_format = '#,##0'; c.alignment = Alignment(horizontal='right')
    for j in (1, 2, 3):
        ws.cell(r, j).border = Border(bottom=line)
        if fill: ws.cell(r, j).fill = PatternFill('solid', fgColor=fill)
    r += 1
    return r - 1
ws.cell(r, 1, 'I. THU NHẬP').font = Font(name='Times New Roman', bold=True, color='1F4E5F'); r += 1
inc = [money_row(l, f'[{l}]') for l in ['Lương cơ bản', 'Lương hiệu quả', 'Làm thêm giờ', 'Phụ cấp ăn ca', 'Phụ cấp điện thoại', 'Công tác phí']]
tot = money_row('Tổng thu nhập', f'=SUM(C{inc[0]}:C{inc[-1]})', True, 'EEF3F5'); r += 1
ws.cell(r, 1, 'II. KHẤU TRỪ').font = Font(name='Times New Roman', bold=True, color='1F4E5F'); r += 1
ded = [money_row(l, f'[{k}]') for l, k in [('BHXH (8%)', 'BHXH'), ('BHYT (1,5%)', 'BHYT'), ('BHTN (1%)', 'BHTN'), ('Thuế TNCN', 'Thuế TNCN'), ('Tạm ứng', 'Tạm ứng')]]
kt = money_row('Tổng khấu trừ', f'=SUM(C{ded[0]}:C{ded[-1]})', True, 'EEF3F5'); r += 1
ws.cell(r, 2, 'THỰC LĨNH').font = Font(name='Times New Roman', bold=True, size=13, color='FFFFFF')
c = ws.cell(r, 3, f'=C{tot}-C{kt}'); c.font = Font(name='Times New Roman', bold=True, size=13, color='FFE2A8'); c.number_format = '#,##0'; c.alignment = Alignment(horizontal='right')
for j in (1, 2, 3): ws.cell(r, j).fill = PatternFill('solid', fgColor='1F4E5F')
r += 1
ws.cell(r, 1, 'Bằng chữ: [Bằng chữ]').font = Font(name='Times New Roman', italic=True, size=10); r += 1
ws.cell(r, 1, 'Ghi chú: [Ghi chú]').font = Font(name='Times New Roman', italic=True, size=10, color='6B777D')
ws.page_setup.paperSize = 9; ws.print_options.horizontalCentered = True
wb.save(f'{OUT}/mau-phieu-luong.xlsx')
for x in rows[:2]: print(x)
